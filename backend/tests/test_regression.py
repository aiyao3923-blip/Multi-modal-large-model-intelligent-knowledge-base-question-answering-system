"""隔离回归：仅在本任务生成的 SQLite/上传目录中写入，不触碰业务数据库。"""
from __future__ import annotations
import json
import os
import shutil
import sys
import tempfile
import time
import unittest
from datetime import datetime, timedelta, timezone
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[2]
RUNS = ROOT / 'storage' / 'refactor-validation' / 'test-runs'
RUNS.mkdir(parents=True, exist_ok=True)
TEMP = Path(tempfile.mkdtemp(prefix='regression-', dir=RUNS)).resolve()
assert TEMP.is_relative_to(RUNS.resolve())
CONFIG = TEMP / 'config.yaml'
CONFIG.write_text('app:\n  name: 隔离回归测试\n  version: V1.0\nstorage:\n  data_dir: "' + (TEMP / 'data').as_posix() + '"\n  upload_dir: "' + (TEMP / 'uploads').as_posix() + '"\n  vector_dir: "' + (TEMP / 'vectors').as_posix() + '"\nsecurity:\n  secret_key: isolated-regression-only-not-production\n  pbkdf2_iterations: 1000\n', encoding='utf-8')
os.environ['APP_CONFIG'] = str(CONFIG)
os.environ.pop('APP_EMBEDDING_MODEL_DIR', None)
os.environ.pop('APP_OCR_ENGINE', None)
sys.path.insert(0, str(ROOT / 'backend'))
from fastapi.testclient import TestClient
from sqlalchemy import event, func, select
from app.main import app
from app.core.config import settings
from app.core.security import hash_password
from app.db.session import Base, get_engine, get_session_factory
from app.db.models import Role, User, KnowledgeBase, UserKbAccess, Document, Chunk, ParseTask, Citation, Message, QaFeedback, ChatSession, StatEvent, LlmChannel, IndustryTemplate
from app.engines.keyword_index import KeywordIndex, _segment
from app.engines.doc_ingest import ingest_document
from app.engines.vector_store import VectorStore, cosine
from app.workers.tasks import run_sync_parse
from app.services import qa_service

PASSWORD = 'OnlyTest2026!'
TEXT = '办理居住证需要本人身份证、居住证明、近期一寸照片。申请材料提交后在五个工作日内办结。'
METRICS = {}

class Regression(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        assert settings.db_path.resolve().is_relative_to(TEMP)
        cls.engine = get_engine()
        cls.factory = get_session_factory()
        cls.client = TestClient(app)
        cls.worker_patch = patch('app.api.kb.submit_parse', side_effect=run_sync_parse)
        cls.worker_patch.start()

    @classmethod
    def tearDownClass(cls):
        cls.worker_patch.stop()
        cls.client.close()
        cls.engine.dispose()

    def setUp(self):
        assert Path(self.engine.url.database).resolve().is_relative_to(TEMP)
        Base.metadata.drop_all(self.engine)
        Base.metadata.create_all(self.engine)
        KeywordIndex.invalidate()
        with self.factory() as db:
            db.add_all([Role(code=c,name=c) for c in ('admin','operator','user')])
            salt, digest = hash_password(PASSWORD)
            db.add_all([User(username=n,display_name=n,role_code=r,password_salt=salt,password_digest=digest,is_active=True) for n,r in (('admin','admin'),('user1','user'),('user2','user'),('operator','operator'))])
            db.commit()
            users = {u.username:u.id for u in db.scalars(select(User))}
            self.users = users
            kb = KnowledgeBase(name='政务服务知识库',owner_id=users['admin'],status='enabled',description='办事指南')
            private = KnowledgeBase(name='管理员专有知识库',owner_id=users['admin'],status='enabled')
            db.add_all([kb,private]);db.flush()
            self.kb, self.private_kb = kb.id, private.id
            db.add_all([UserKbAccess(kb_id=kb.id,role_code=r) for r in ('user','operator')])
            db.commit()
        self.headers = {}
        for username in self.users:
            result = self.client.post('/api/auth/login',json={'username':username,'password':PASSWORD})
            self.assertEqual(result.status_code,200,result.text)
            self.headers[username] = {'Authorization':'Bearer '+result.json()['token']}

    def req(self, method, url, who='admin', **kwargs):
        return self.client.request(method,url,headers=self.headers[who],**kwargs)

    def upload(self, text=TEXT, name='办事指南.md', kb=None):
        response = self.req('POST',f'/api/kbs/{kb or self.kb}/docs',files={'file':(name,text.encode('utf-8'),'text/plain')})
        self.assertEqual(response.status_code,200,response.text)
        return response.json()['id']

    def ask(self, question='办理居住证需要哪些材料？', who='user1', kb=None, session=None):
        result = self.req('POST','/api/qa/ask',who=who,json={'kb_id':kb or self.kb,'question':question,'session_id':session})
        self.assertEqual(result.status_code,200,result.text)
        return result.json()

    def test_health_and_auth_boundaries(self):
        self.assertEqual(self.client.get('/api/health').status_code,200)
        self.assertEqual(self.client.get('/api/kbs').status_code,401)
        self.assertEqual(self.client.post('/api/auth/login',json={'username':'admin','password':'wrong'}).status_code,400)
        for path in ('/api/sys/users','/api/stats/overview'):
            self.assertEqual(self.req('GET',path,who='user1').status_code,403)
        self.assertEqual(self.req('GET','/api/tasks/page',who='user1').status_code,200)
        self.assertEqual(self.req('GET','/api/auth/me',who='user1').json()['username'],'user1')
        self.assertEqual(self.req('GET','/api/qa/capabilities',who='user1').json()['mode'],'degraded')

    def test_knowledge_creation_validation_and_owner_access(self):
        response = self.req('POST','/api/kbs',who='operator',json={'name':'  运营员新库  ','role_codes':[]})
        self.assertEqual(response.status_code,200,response.text)
        kb_id = response.json()['id']
        self.assertEqual(response.json()['name'],'运营员新库')
        self.assertEqual(self.req('GET',f'/api/kbs/{kb_id}',who='operator').status_code,200)
        self.assertEqual(self.req('GET',f'/api/kbs/{kb_id}',who='user1').status_code,403)
        self.assertEqual(self.req('POST','/api/kbs',json={'name':'   '}).status_code,422)
        self.assertEqual(self.req('POST','/api/kbs',json={'name':'x','role_codes':['unknown']}).status_code,422)
        self.assertEqual(self.req('PATCH',f'/api/kbs/{kb_id}',json={'status':'broken'}).status_code,422)
        self.assertEqual(self.req('POST',f'/api/kbs/{kb_id}/grant',json={'role_codes':['user','user']}).status_code,403)
        self.assertEqual(self.req('POST',f'/api/kbs/{kb_id}/grant',who='operator',json={'role_codes':['user','user']}).status_code,200)
        self.assertEqual(self.req('GET',f'/api/kbs/{kb_id}',who='user1').json()['granted_roles'],['user'])

    def test_upload_parse_search_pagination_and_duplicate(self):
        doc = self.upload()
        other = self.upload('实验室设备维护周期为三个月。','设备维护.txt')
        result = self.req('GET',f'/api/kbs/{self.kb}/docs',params={'q':'办事','status':'success','size':1}).json()
        self.assertEqual(result['total'],1)
        self.assertEqual(result['items'][0]['id'],doc)
        self.assertGreater(result['items'][0]['chunk_count'],0)
        chunks = self.req('GET',f'/api/docs/{doc}/chunks').json()
        self.assertIn('居住证',chunks[0]['content'])
        json.loads(chunks[0]['meta_json'])
        self.assertEqual(self.req('POST',f'/api/kbs/{self.kb}/docs',files={'file':('重复.md',TEXT.encode())}).status_code,409)
        self.assertEqual(self.req('POST',f'/api/kbs/{self.kb}/docs',files={'file':('空.txt',b'')}).status_code,400)
        self.assertEqual(self.req('POST',f'/api/kbs/{self.kb}/docs',files={'file':('old.xls',b'x')}).status_code,400)
        self.assertEqual(self.req('GET',f'/api/kbs/{self.kb}/docs',params={'size':101}).status_code,422)
        self.assertEqual(self.req('GET',f'/api/kbs/{self.kb}/docs',params={'q':'%'}).json()['total'],0)

    def test_csv_parser_retains_headers_and_row_values(self):
        doc = self.upload('事项,材料,时限\n居住证,身份证与居住证明,五个工作日\n', '事项清单.csv')
        chunks = self.req('GET',f'/api/docs/{doc}/chunks').json()
        self.assertTrue(any('事项：居住证' in c['content'] and '材料：身份证' in c['content'] for c in chunks))
        self.assertTrue(self.ask()['citations'])

    def test_real_answer_citations_history_and_feedback(self):
        doc = self.upload()
        answer = self.ask()
        self.assertEqual(answer['mode'],'degraded')
        self.assertTrue(answer['citations'])
        self.assertEqual(answer['citations'][0]['doc_id'],doc)
        self.assertIn('身份证', ''.join(c['content_excerpt'] for c in answer['citations']))
        self.assertNotIn('身份证', answer['answer'])
        self.assertEqual(answer['generation_status'], 'unconfigured')
        history = self.req('GET',f'/api/qa/sessions/{answer["session_id"]}/messages',who='user1')
        self.assertEqual(history.status_code,200,history.text)
        self.assertEqual([m['role'] for m in history.json()],['user','assistant'])
        self.assertEqual(history.json()[1]['citations'],answer['citations'])
        fb = self.req('POST','/api/qa/feedback',who='user1',json={'message_id':answer['message_id'],'feedback_type':'dislike','comment':'需要补充材料细节'})
        self.assertEqual(fb.status_code,200,fb.text)
        fid = fb.json()['id']
        pending = self.req('GET','/api/qa/feedbacks/page',params={'status':'pending','q':'居住证'}).json()
        self.assertEqual(pending['total'],1)
        self.assertEqual(pending['items'][0]['answer'], answer['answer'])
        self.assertNotIn('身份证', pending['items'][0]['answer'])
        self.assertEqual(self.req('POST',f'/api/qa/feedbacks/{fid}/handle',params={'action':'handled','note':'已核验办事指南'}).json()['status'],'handled')
        self.assertEqual(self.req('POST',f'/api/qa/feedbacks/{fid}/handle',params={'action':'pending'}).json()['status'],'pending')
        self.req('POST','/api/qa/feedback',who='user1',json={'message_id':answer['message_id'],'feedback_type':'like'})
        history = self.req('GET',f'/api/qa/sessions/{answer["session_id"]}/messages',who='user1').json()
        self.assertEqual(history[1]['feedback_type'],'like')
        with self.factory() as db:
            self.assertEqual(db.scalar(select(func.count(QaFeedback.id))),1)
            self.assertEqual(db.get(QaFeedback,fid).handler_note,'')

    def test_cross_account_and_library_isolation(self):
        self.upload()
        answer = self.ask()
        sid, mid = answer['session_id'], answer['message_id']
        self.assertEqual(self.req('GET',f'/api/qa/sessions/{sid}/messages',who='user2').status_code,404)
        self.assertEqual(self.req('DELETE',f'/api/qa/sessions/{sid}',who='user2').status_code,404)
        self.assertEqual(self.req('POST','/api/qa/feedback',who='user2',json={'message_id':mid,'feedback_type':'like'}).status_code,404)
        self.assertEqual(self.req('POST','/api/qa/ask',who='user1',json={'kb_id':self.private_kb,'question':'秘密'}).status_code,403)
        self.assertEqual(self.req('GET','/api/tasks/page',who='operator',params={'kb_id':self.private_kb}).status_code,403)
        self.assertEqual(self.req('GET','/api/stats/trend',who='operator',params={'kb_id':self.private_kb}).status_code,403)
        self.assertNotIn(self.private_kb,[k['id'] for k in self.req('GET','/api/kbs',who='operator').json()])

    def test_disabled_and_revoked_knowledge(self):
        self.upload()
        answer = self.ask()
        self.req('PATCH',f'/api/kbs/{self.kb}',json={'status':'disabled'})
        self.assertEqual(self.req('POST','/api/qa/ask',who='user1',json={'kb_id':self.kb,'question':'居住证'}).status_code,409)
        self.assertEqual(self.req('GET',f'/api/qa/sessions/{answer["session_id"]}/messages',who='user1').status_code,200)
        self.req('POST',f'/api/kbs/{self.kb}/grant',json={'role_codes':[]})
        self.assertEqual(self.req('GET',f'/api/qa/sessions/{answer["session_id"]}/messages',who='user1').status_code,403)
        self.assertEqual(self.req('GET','/api/qa/sessions',who='user1').json(),[])

    def test_sessions_rename_boundaries_and_cleanup(self):
        self.upload()
        answer = self.ask()
        sid = answer['session_id']
        self.req('POST','/api/qa/feedback',who='user1',json={'message_id':answer['message_id'],'feedback_type':'like'})
        self.assertEqual(self.req('PATCH',f'/api/qa/sessions/{sid}',who='user1',params={'title':'   '}).status_code,422)
        self.assertEqual(self.req('PATCH',f'/api/qa/sessions/{sid}',who='user1',params={'title':'材料办理指南'}).json()['title'],'材料办理指南')
        self.assertEqual(self.req('GET','/api/qa/sessions',who='user1',params={'q':'材料'}).json()[0]['id'],sid)
        self.assertEqual(self.req('POST','/api/qa/ask',who='admin',json={'kb_id':self.private_kb,'question':'hello','session_id':sid}).status_code,404)
        self.assertEqual(self.req('DELETE',f'/api/qa/sessions/{sid}',who='user1').status_code,200)
        with self.factory() as db:
            for model in (ChatSession,Message,Citation,QaFeedback):
                self.assertEqual(db.scalar(select(func.count()).select_from(model)),0)
            self.assertEqual(db.scalar(select(StatEvent)).session_id,None)
        self.assertEqual(self.req('GET','/api/stats/overview').json()['qa_total'],1)

    def test_reparse_preserves_citation_snapshot_and_delete_invalidates_index(self):
        doc_id = self.upload()
        answer = self.ask()
        original = answer['citations'][0]['content_excerpt']
        with self.factory() as db:
            doc = db.get(Document,doc_id)
            (settings.root_dir / doc.storage_path).write_text('新的业务：图书借阅需要读者证。',encoding='utf-8')
        self.assertEqual(self.req('POST',f'/api/kbs/{self.kb}/docs/{doc_id}/reparse').status_code,200)
        history = self.req('GET',f'/api/qa/sessions/{answer["session_id"]}/messages',who='user1').json()
        self.assertEqual(history[1]['citations'][0]['content_excerpt'],original)
        self.assertIsNone(history[1]['citations'][0]['chunk_id'])
        self.assertTrue(KeywordIndex().search('读者证',self.kb))
        self.assertEqual(KeywordIndex().search('居住证',self.kb),[])
        self.assertEqual(self.req('DELETE',f'/api/kbs/{self.kb}/docs/{doc_id}').status_code,200)
        self.assertEqual(KeywordIndex().search('读者证',self.kb),[])
        self.assertEqual(self.req('GET',f'/api/qa/sessions/{answer["session_id"]}/messages',who='user1').status_code,200)

    def test_failures_reach_terminal_state_and_active_tasks_are_guarded(self):
        with patch('app.engines.doc_ingest._run_extract',side_effect=RuntimeError('synthetic failure')):
            doc = self.upload('损坏文件模拟','异常文档.txt')
        tasks = self.req('GET','/api/tasks/page').json()['items']
        self.assertEqual(tasks[0]['phase'],'failed')
        self.assertIsNotNone(tasks[0]['started_at'])
        self.assertIsNotNone(tasks[0]['finished_at'])
        self.assertTrue(tasks[0]['error_message'])
        with patch('app.api.kb.submit_parse'):
            self.assertEqual(self.req('POST',f'/api/kbs/{self.kb}/docs/{doc}/reparse').status_code,200)
            self.assertEqual(self.req('POST',f'/api/kbs/{self.kb}/docs/{doc}/reparse').status_code,409)
            self.assertEqual(self.req('DELETE',f'/api/kbs/{self.kb}/docs/{doc}').status_code,409)
        self.assertEqual(self.req('GET',f'/api/kbs/{self.kb}/stats').json()['pending_docs'],1)

    def test_whitespace_and_unrelated_queries_never_call_models(self):
        self.upload()
        self.assertEqual(self.req('POST','/api/qa/ask',who='user1',json={'kb_id':self.kb,'question':'  '}).status_code,422)
        with patch('app.services.qa_service._call_llm_safe') as mocked:
            response = self.ask('zzzxylophoneunknown')
            self.assertEqual(response['mode'],'none')
            self.assertEqual(response['citations'],[])
            mocked.assert_not_called()
        self.assertEqual(KeywordIndex().search('？？？',self.kb),[])

    def test_model_fallback_uses_real_successful_mode_without_network(self):
        self.upload()
        with self.factory() as db:
            db.add_all([LlmChannel(name=t,channel_type=t,model_name='fixture',base_url='http://127.0.0.1:9',enabled=True,health_status='unknown') for t in ('local','online')])
            db.commit()
        with patch('app.services.qa_service._call_llm_safe',side_effect=[None,'请准备身份证与居住证明。[来源1]']) as mocked:
            response = self.ask()
            self.assertEqual(mocked.call_count,2)
            self.assertEqual(response['mode'],'online')
            self.assertFalse(response['degraded'])
        self.assertEqual(self.req('GET','/api/stats/overview').json()['mode_dist']['online'],1)
        point = self.req('GET','/api/stats/trend').json()[-1]
        self.assertEqual(point['online'],1)
        self.assertEqual(point['local'],0)
        self.assertEqual(point['cited'],1)

    def test_statistics_scope_window_and_empty_days(self):
        self.upload()
        self.ask()
        self.upload('管理员内部预算资料','私有资料.txt',kb=self.private_kb)
        self.ask('预算',who='admin',kb=self.private_kb)
        admin = self.req('GET','/api/stats/overview').json()
        operator = self.req('GET','/api/stats/overview',who='operator').json()
        self.assertEqual(admin['qa_total'],2)
        self.assertEqual(operator['qa_total'],1)
        self.assertEqual(operator['doc_total'],1)
        self.assertEqual(operator['cite_rate'],100)
        with self.factory() as db:
            db.add(StatEvent(event_type='qa_completed',kb_id=self.kb,mode='local',created_at=datetime.now(timezone.utc).replace(tzinfo=None)-timedelta(days=60)))
            db.commit()
        trend = self.req('GET','/api/stats/trend',who='operator').json()
        self.assertEqual(len(trend),30)
        self.assertEqual(sum(p['qa_count'] for p in trend),1)
        self.assertEqual(sum(p['local'] for p in trend),0)
        self.assertTrue(any(p['qa_count']==0 for p in trend))

    def test_user_validation_and_self_lockout_guard(self):
        uid = self.users['admin']
        self.assertEqual(self.req('PATCH',f'/api/sys/users/{uid}',json={'is_active':False}).status_code,409)
        self.assertEqual(self.req('PATCH',f'/api/sys/users/{uid}',json={'role_code':'user'}).status_code,409)
        self.assertEqual(self.req('POST','/api/sys/users',json={'username':'newuser','password':'123','role_code':'user'}).status_code,422)
        self.assertEqual(self.req('POST','/api/sys/users',json={'username':'newuser','password':PASSWORD,'role_code':'unknown'}).status_code,422)
        result = self.req('POST','/api/sys/users',json={'username':'newuser','password':PASSWORD,'role_code':'operator'})
        self.assertEqual(result.status_code,200,result.text)
        self.assertNotIn('password',result.text)
        self.req('PATCH',f'/api/sys/users/{self.users["user2"]}',json={'is_active':False})
        self.assertEqual(self.req('GET','/api/auth/me',who='user2').status_code,401)

    def test_channel_validation_and_secret_not_returned_or_redirected(self):
        secret = 'fixture-key-do-not-display-1234'
        response = self.req('POST','/api/sys/channels',json={'name':'测试通道','channel_type':'online','model_name':'fixture','base_url':'https://example.invalid','api_key':secret})
        self.assertEqual(response.status_code,200,response.text)
        cid = response.json()['id']
        self.assertFalse(response.json()['enabled'])
        self.assertNotIn(secret,response.text)
        self.assertNotIn('1234',self.req('GET','/api/sys/channels').text)
        self.assertEqual(self.req('PATCH',f'/api/sys/channels/{cid}',json={'base_url':'https://other.invalid'}).status_code,409)
        self.assertEqual(self.req('PATCH',f'/api/sys/channels/{cid}',json={'timeout_s':1000}).status_code,422)
        self.assertEqual(self.req('POST','/api/sys/channels',json={'name':'x','channel_type':'online','model_name':'x','base_url':'file:///etc/passwd'}).status_code,422)

    def count_selects(self, call):
        queries=[]
        def record(_conn,_cursor,statement,_params,_context,_many):
            if statement.lstrip().upper().startswith('SELECT'):
                queries.append(statement)
        event.listen(self.engine,'before_cursor_execute',record)
        try:
            response=call()
            self.assertEqual(response.status_code,200,response.text)
        finally:
            event.remove(self.engine,'before_cursor_execute',record)
        return len(queries)

    def test_authenticated_user_row_is_reused_only_per_request(self):
        queries = []
        def record(_conn, _cursor, statement, _params, _context, _many):
            if statement.lstrip().upper().startswith('SELECT') and 'FROM users' in statement:
                queries.append(statement)
        event.listen(self.engine, 'before_cursor_execute', record)
        try:
            self.assertEqual(self.req('GET', '/api/kbs', who='user1').status_code, 200)
        finally:
            event.remove(self.engine, 'before_cursor_execute', record)
        self.assertEqual(len(queries), 1, 'Reuse the authenticated row within one request, not across requests')
        with self.factory() as db:
            db.get(User, self.users['user1']).is_active = False
            db.commit()
        self.assertEqual(self.req('GET', '/api/kbs', who='user1').status_code, 401)
        with self.factory() as db:
            user = db.get(User, self.users['user1'])
            user.is_active, user.role_code = True, 'admin'
            db.commit()
        self.assertEqual(self.req('GET', '/api/sys/users', who='user1').status_code, 200)
        with self.factory() as db:
            db.get(User, self.users['user1']).role_code = 'user'
            db.commit()
        self.assertEqual(self.req('GET', '/api/sys/users', who='user1').status_code, 403)

    def test_list_query_budget_with_assigned_templates(self):
        with self.factory() as db:
            template = IndustryTemplate(code='query-budget', name='真实关联模板')
            db.add(template)
            db.flush()
            db.get(KnowledgeBase, self.kb).template_id = template.id
            db.add_all([KnowledgeBase(name=f'带模板知识库{i}', owner_id=self.users['admin'], template_id=template.id) for i in range(30)])
            db.commit()
        count = self.count_selects(lambda: self.req('GET', '/api/kbs'))
        self.assertLessEqual(count, 10)
        rows = self.req('GET', '/api/kbs').json()
        self.assertEqual(next(row for row in rows if row['id'] == self.kb)['template_name'], '真实关联模板')

    def test_list_query_counts_do_not_grow_with_libraries_or_documents(self):
        small=self.count_selects(lambda:self.req('GET','/api/kbs'))
        with self.factory() as db:
            db.add_all([KnowledgeBase(name=f'额外知识库{i}',owner_id=self.users['admin']) for i in range(30)])
            db.commit()
        large=self.count_selects(lambda:self.req('GET','/api/kbs'))
        self.assertLessEqual(large,small+1)
        self.assertLessEqual(large,10)
        with self.factory() as db:
            db.add_all([Document(kb_id=self.kb,filename=f'doc-{i}.txt',file_type='txt',uploaded_by=self.users['admin'],status='success') for i in range(50)])
            db.commit()
        few=self.count_selects(lambda:self.req('GET',f'/api/kbs/{self.kb}/docs',params={'size':1}))
        many=self.count_selects(lambda:self.req('GET',f'/api/kbs/{self.kb}/docs',params={'size':100}))
        self.assertLessEqual(many,few+1)
        METRICS['list_selects']={'two_kbs':small,'32_kbs':large,'one_doc':few,'50_docs':many}

    def test_cached_bm25_reuses_corpus_tokens_and_benchmark(self):
        with self.factory() as db:
            doc=Document(kb_id=self.kb,filename='性能测试.txt',file_type='txt',status='success',uploaded_by=self.users['admin'])
            db.add(doc);db.flush()
            db.add_all([Chunk(kb_id=self.kb,document_id=doc.id,seq_no=i+1,content=f'事项{i}：居住证办理需要身份证、居住证明。材料齐全后五个工作日办结。') for i in range(1500)])
            db.commit()
        index=KeywordIndex()
        with patch('app.engines.keyword_index._segment',wraps=_segment) as tokenize:
            start=time.perf_counter();cold=index.search('居住证材料',self.kb);cold_ms=(time.perf_counter()-start)*1000
            before=tokenize.call_count
            times=[]
            for _ in range(12):
                start=time.perf_counter();warm=index.search('居住证材料',self.kb);times.append((time.perf_counter()-start)*1000)
            self.assertEqual(tokenize.call_count-before,12)
            self.assertEqual(cold,warm)
        average=sum(times)/len(times)
        self.assertLess(average,cold_ms)
        METRICS['bm25']={'chunks':1500,'cold_ms':round(cold_ms,3),'warm_mean_ms':round(average,3),'warm_samples':len(times),'tokenizations_per_warm_query':1}

    def test_vectors_not_limited_to_first_2000_and_dimension_guard(self):
        with self.factory() as db:
            doc=Document(kb_id=self.kb,filename='vector-fixture.txt',file_type='txt',status='success',uploaded_by=self.users['admin'])
            db.add(doc);db.flush()
            db.add_all([Chunk(kb_id=self.kb,document_id=doc.id,seq_no=i+1,content=str(i),vector_json='[0,1]') for i in range(2001)])
            target=Chunk(kb_id=self.kb,document_id=doc.id,seq_no=2002,content='目标向量',vector_json='[1,0]')
            db.add(target);db.commit();target_id=target.id
        self.assertEqual(VectorStore().search([1,0],self.kb,top_k=1,threshold=.9)[0]['chunk_id'],target_id)
        self.assertEqual(cosine([1,0],[1]),0)

    def upload_bytes(self, payload, filename, mime):
        response = self.req('POST', f'/api/kbs/{self.kb}/docs', files={'file': (filename, payload, mime)})
        self.assertEqual(response.status_code, 200, response.text)
        doc_id = response.json()['id']
        chunks = self.req('GET', f'/api/docs/{doc_id}/chunks')
        self.assertEqual(chunks.status_code, 200, chunks.text)
        return doc_id, chunks.json()

    def test_real_docx_paragraph_table_and_citation(self):
        from io import BytesIO
        from docx import Document as WordDocument
        word = WordDocument()
        word.add_heading('居住证 Word 办理指南', 0)
        word.add_paragraph(TEXT)
        table = word.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = '材料', '要求'
        table.cell(1, 0).text, table.cell(1, 1).text = '身份证', '原件与复印件'
        output = BytesIO()
        word.save(output)
        doc_id, chunks = self.upload_bytes(output.getvalue(), 'Word办理指南.docx', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')
        self.assertTrue(any('原件与复印件' in c['content'] and json.loads(c['meta_json']).get('rows') == 2 for c in chunks))
        answer = self.ask()
        self.assertTrue(any(c['doc_id'] == doc_id for c in answer['citations']))
        self.assertIn('身份证', ''.join(c['content_excerpt'] for c in answer['citations']))
        self.assertIn('未启用生成模型', answer['answer'])

    def test_real_xlsx_sheets_and_table_values(self):
        from io import BytesIO
        from openpyxl import Workbook
        workbook = Workbook()
        first = workbook.active
        first.title = '办理材料'
        first.append(['事项', '材料', '时限'])
        first.append(['居住证', '身份证与居住证明', '五个工作日'])
        second = workbook.create_sheet('咨询时间')
        second.append(['星期', '开放时间'])
        second.append(['周一至周五', '09:00-17:00'])
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        doc_id, chunks = self.upload_bytes(output.getvalue(), '材料清单.xlsx', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertEqual({json.loads(c['meta_json']).get('sheet') for c in chunks}, {'办理材料', '咨询时间'})
        self.assertTrue(any('居住证明' in c['content'] for c in chunks))
        self.assertTrue(any(c['doc_id'] == doc_id for c in self.ask()['citations']))

    def test_real_pdf_text_layers_keep_page_numbers(self):
        from io import BytesIO
        from pypdf import PdfWriter
        from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
        writer = PdfWriter()
        for text in ('Residence permit requires passport and address proof.', 'Application status is available using receipt number.'):
            page = writer.add_blank_page(width=400, height=300)
            font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
            page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
            stream = DecodedStreamObject()
            stream.set_data(('BT /F1 12 Tf 30 250 Td (' + text + ') Tj ET').encode('ascii'))
            page[NameObject('/Contents')] = writer._add_object(stream)
        output = BytesIO()
        writer.write(output)
        self.assertTrue(output.getvalue().startswith(b'%PDF-'))
        doc_id, chunks = self.upload_bytes(output.getvalue(), 'permit-guide.pdf', 'application/pdf')
        self.assertEqual([c['page_no'] for c in chunks], [1, 2])
        self.assertIn('passport', chunks[0]['content'])
        self.assertIn('receipt number', chunks[1]['content'])
        answer = self.ask(question='passport address proof')
        self.assertTrue(any(c['doc_id'] == doc_id and c['page_no'] == 1 for c in answer['citations']))

    def test_html_keeps_div_content_and_document_order_without_script(self):
        html = '<html><head><script>unsafe()</script></head><body><!--hidden-not-knowledge--><h1>办理指南</h1><div>居住证需要<span>身份证和居住证明</span></div><table><tr><th>时间</th></tr><tr><td>五个工作日</td></tr></table><p>请保留申请回执</p></body></html>'
        doc_id = self.upload(html, '网页指南.html')
        chunks = self.req('GET', f'/api/docs/{doc_id}/chunks').json()
        combined = chr(10).join(c['content'] for c in chunks)
        self.assertIn('身份证和居住证明', combined)
        self.assertNotIn('unsafe()', combined)
        self.assertNotIn('hidden-not-knowledge', combined)
        self.assertLess(combined.index('办理指南'), combined.index('身份证'))
        self.assertLess(combined.index('身份证'), combined.index('五个工作日'))
        self.assertLess(combined.index('五个工作日'), combined.index('请保留'))

    def test_image_ocr_is_explicit_local_and_does_not_download(self):
        from app.engines.doc_parser import ParseError
        from app.engines.doc_parser.image_parser import ImageParser, _try_paddle
        with patch.dict(os.environ, {'APP_OCR_ENGINE': 'none'}), patch('app.engines.doc_parser.image_parser._try_paddle') as paddle, patch('app.engines.doc_parser.image_parser._try_tesseract') as tesseract:
            with self.assertRaisesRegex(ParseError, '未启用'):
                ImageParser().parse(TEMP / 'image.png')
            paddle.assert_not_called()
            tesseract.assert_not_called()
        with patch.dict(os.environ, {'APP_PADDLE_DET_MODEL_DIR': '', 'APP_PADDLE_REC_MODEL_DIR': '', 'APP_PADDLE_CLS_MODEL_DIR': ''}):
            self.assertIsNone(_try_paddle(TEMP / 'image.png'))
        with patch.dict(os.environ, {'APP_OCR_ENGINE': 'tesseract'}), patch('app.engines.doc_parser.image_parser._try_tesseract', return_value='本地图片识别文本'):
            part = ImageParser().parse(TEMP / 'image.png')[0]
            self.assertEqual(part.text, '本地图片识别文本')
            self.assertEqual(part.meta['engine'], 'tesseract')

    def test_chunking_retains_heading_only_documents(self):
        from app.engines.chunker import chunk_parts
        from app.engines.doc_parser import Part
        chunks = chunk_parts([Part(text='一级标题', kind='heading'), Part(text='二级标题', kind='heading')])
        self.assertEqual(len(chunks), 1)
        self.assertIn('一级标题', chunks[0]['content'])
        self.assertIn('二级标题', chunks[0]['content'])

if __name__=='__main__':
    started=time.perf_counter()
    result=unittest.TextTestRunner(verbosity=2).run(unittest.defaultTestLoader.loadTestsFromTestCase(Regression))
    output={'passed':result.wasSuccessful(),'tests':result.testsRun,'failures':len(result.failures),'errors':len(result.errors),'duration_seconds':round(time.perf_counter()-started,2),'metrics':METRICS,'database_isolated':True}
    (ROOT/'storage'/'refactor-validation'/'backend-results.json').write_text(json.dumps(output,ensure_ascii=False,indent=2),encoding='utf-8')
    # 仅移除本进程新建且经绝对路径核验的隔离测试目录。
    assert TEMP.is_relative_to(RUNS.resolve()) and TEMP.name.startswith('regression-')
    shutil.rmtree(TEMP)
    raise SystemExit(0 if result.wasSuccessful() else 1)
